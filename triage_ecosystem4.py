import os
import pandas as pd
import numpy as np
import re
import datetime
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

# File Input Environment Definitions
MASTER_INPUT = "Cayuse-Oracle award True Up - jth6-Copy.xlsx"
PRIME_ASST_FILE = "Assistance_PrimeAwardSummaries_2026-06-25_H19M12S13_1.csv"
PRIME_CONT_FILE = "Contracts_PrimeAwardSummaries_2026-06-25_H19M12S08_1.csv"
SUB_ASST_FILE = "Assistance_Subawards_2026-06-25_H19M12S41_1.csv"
SUB_CONT_FILE = "Contracts_Subawards_2026-06-25_H19M12S29_1.csv"

# Current data baseline extraction execution runtime anchor date
PIPELINE_RUN_DATE = pd.to_datetime("2026-06-26").date()

print("🚀 Launching Corrected Three-Way Locked Triage Engine...")

def advanced_clean_id(text):
    if pd.isna(text): return ""
    text = str(text).upper().strip()
    text = re.sub(r'^(NIH|NSF|DOD|DOHA|ONR|DOE|USDA|DHHS)\s+', '', text)
    text = re.sub(r'[^A-Z0-9-]', '', text)
    text = re.sub(r'-[A-Z0-9]{1,2}$', '', text) 
    text = re.sub(r'[^A-Z0-9]', '', text)
    if re.match(r'^[A-Z]+\d{7}$', text): text = re.sub(r'^[A-Z]+', '', text)
    if re.match(r'^\d[A-Z]\d{2}', text): text = text[1:]
    return text

def tokenize_string(s):
    if pd.isna(s): return []
    tokens = re.findall(r'[A-Z0-9\-]{4,}', str(s).upper())
    noise = {'RICE', 'WEST', 'GRANT', 'FEDERAL', 'CONTR', 'CONTRACT', 'AMENDMENT', 'MODIFICATION', 'BLANK', 'NONE', 'FOUNDATION', 'UNIVERSITY'}
    # 🧠 FIX: Changed 'w' to 't' to match the list comprehension iterator token
    return [t for t in tokens if t not in noise]

def parse_mixed_date(val):
    if pd.isna(val): return None
    val_str = str(val).strip()
    if val_str.replace('.0', '').isdigit():
        clean_int = int(val_str.replace('.0', ''))
        return pd.to_datetime(clean_int, unit='D', origin='1899-12-30').date()
    return pd.to_datetime(val_str, errors='coerce').date()

def save_as_structured_table(writer, dataframe, name_of_sheet):
    dataframe.to_excel(writer, sheet_name=name_of_sheet, index=False)
    worksheet = writer.sheets[name_of_sheet]
    total_rows = len(dataframe) + 1
    total_cols = len(dataframe.columns) if len(dataframe) > 0 else 1
    column_letter = get_column_letter(total_cols)
    cell_range = f"A1:{column_letter}{total_rows}"
    table_id = "".join(c for c in name_of_sheet if c.isalnum())
    excel_table = Table(displayName=table_id, ref=cell_range)
    excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    worksheet.add_table(excel_table)

# Ingest source matrices
master_df = pd.read_excel(MASTER_INPUT, sheet_name="Intermediate DATA")
gsum_df = pd.read_excel(MASTER_INPUT, sheet_name="Grant Summary")

gsum_df['clean_id'] = gsum_df['AWARD_NUMBER_SCRUB'].astype(str).str.upper().str.strip()
type_map = dict(zip(gsum_df['clean_id'], gsum_df['AWARD_TYPE_NAME']))
purpose_map = dict(zip(gsum_df['clean_id'], gsum_df['AWARD_PURPOSE']))
mech_map = dict(zip(gsum_df['clean_id'], gsum_df['FUNDING_MECHANISM']))
status_map = dict(zip(gsum_df['clean_id'], gsum_df['AWARD_STATUS']))
pi_map = dict(zip(gsum_df['clean_id'], gsum_df['PI_NAME']))
flow_through_map = dict(zip(gsum_df['clean_id'], gsum_df['FLOW_THROUGH_SPONSOR']))
header_limit_map = dict(zip(gsum_df['clean_id'], gsum_df['HARD_LIMIT_AMOUNT']))
project_budget_map = dict(zip(gsum_df['clean_id'], gsum_df['BUDGET_BRDND_COST']))
start_date_map = dict(zip(gsum_df['clean_id'], gsum_df['AWARD_START_DATE']))
loc_number_map = dict(zip(gsum_df['clean_id'], gsum_df['LOC_NUMBER']))

pa_df = pd.read_csv(PRIME_ASST_FILE, low_memory=False)
pc_df = pd.read_csv(PRIME_CONT_FILE, low_memory=False)
sa_df = pd.read_csv(SUB_ASST_FILE, low_memory=False)
sc_df = pd.read_csv(SUB_CONT_FILE, low_memory=False)

pa_lookup = {advanced_clean_id(r['award_id_fain']): (r['total_obligated_amount'], r['total_funding_amount'], r['period_of_performance_current_end_date'], r['award_id_fain'], "Prime Assistance") for _, r in pa_df.iterrows() if pd.notna(r['award_id_fain'])}
pc_lookup = {advanced_clean_id(r['award_id_piid']): (r['total_obligated_amount'], r['potential_total_value_of_award'], r['period_of_performance_current_end_date'], r['award_id_piid'], "Prime Contract") for _, r in pc_df.iterrows() if pd.notna(r['award_id_piid'])}

sa_lookup = {}
for _, r in sa_df.iterrows():
    val = (r['subaward_amount'], r['subaward_amount'], r['prime_award_period_of_performance_current_end_date'], r['prime_award_fain'], "Subaward Assistance")
    if pd.notna(r['prime_award_fain']): sa_lookup[advanced_clean_id(r['prime_award_fain'])] = val
    if pd.notna(r['subaward_number']): sa_lookup[advanced_clean_id(r['subaward_number'])] = val

sc_lookup = {}
for _, r in sc_df.iterrows():
    val = (r['subaward_amount'], r['subaward_amount'], r['prime_award_period_of_performance_current_end_date'], r['prime_award_piid'], "Subaward Contract")
    if pd.notna(r['prime_award_piid']): sc_lookup[advanced_clean_id(r['prime_award_piid'])] = val
    if pd.notna(r['subaward_number']): sc_lookup[advanced_clean_id(r['subaward_number'])] = val

triaged_records = []

print("🔗 Processing financial variance rules and injecting logic-trace metrics...")
for idx, row in master_df.iterrows():
    fed_id, fed_end_date, fed_obligated, fed_potential, fed_cat = np.nan, np.nan, 0.0, 0.0, "UNMATCHED"
    clean_award_number = str(row['Award Number']).replace('.0', '').upper().strip()
    
    source_status = status_map.get(clean_award_number, row['STATUS'] if 'STATUS' in master_df.columns else "Active")
    status_clean = str(source_status).upper().strip()
    if status_clean == "CLOSED": continue  

    t_verdict, f_verdict, g_verdict = "IN_SYNC", "FINANCIALS_IN_SYNC", "MATCHED_TO_FEDERAL_REGISTRY"
    t_trace, f_trace = "RULE_0: DEFAULT_STEADY_STATE", "RULE_0: DEFAULT_STEADY_STATE"

    primary_id_key = advanced_clean_id(row['Award Number'])
    name_tokens = tokenize_string(row['Award Name'])
    candidates = [primary_id_key] if primary_id_key else []
    candidates.extend([advanced_clean_id(t) for t in name_tokens if t])
    candidates = [c for c in candidates if c]
    
    for k in candidates:
        if k in pa_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = pa_lookup[k]; break
        elif k in pc_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = pc_lookup[k]; break
        elif k in sa_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = sa_lookup[k]; break
        elif k in sc_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = sc_lookup[k]; break

    o_end = parse_mixed_date(row['OR END DATE'])
    c_end = parse_mixed_date(row['CAY END DATE'])
    f_end = parse_mixed_date(fed_end_date)
    
    c_funding = float(row['CAYUSE BUDGET']) if pd.notna(row['CAYUSE BUDGET']) else 0.0
    o_header_limit = float(header_limit_map.get(clean_award_number, c_funding))
    o_project_budget = float(project_budget_map.get(clean_award_number, float(row['OR Allocated Funding']) if pd.notna(row['OR Allocated Funding']) else 0.0))

    source_type = type_map.get(clean_award_number, row['Award Type'] if 'Award Type' in master_df.columns else "Federal")
    source_mechanism = str(mech_map.get(clean_award_number, "Grant")).strip()
    source_flow = flow_through_map.get(clean_award_number, np.nan)
    source_start = parse_mixed_date(start_date_map.get(clean_award_number, None))

    type_str_upper = str(source_type).upper()
    is_direct_sponsor_fed = "FED" in type_str_upper or "FEDERAL" in type_str_upper
    is_flow_through_fed = pd.notna(source_flow) and any(w in str(source_flow).upper() for w in ['NIH', 'NSF', 'FED', 'FEDERAL', 'HEALTH', 'SCIENCE', 'DEFENSE', 'DEPARTMENT', 'AGENCY', 'FOUNDATION'])
    is_federal_funding = is_direct_sponsor_fed or is_flow_through_fed or (fed_cat != "UNMATCHED")

    days_cayuse_minus_fed = (c_end - f_end).days if (c_end and f_end) else np.nan
    delta_oracle_limit_vs_cayuse = abs(o_header_limit - c_funding)
    delta_project_vs_fed_obligated = abs(o_project_budget - fed_obligated)

    # Timeline Engine
    if fed_cat != "UNMATCHED":
        if o_end != f_end or c_end != f_end:
            if f_end and c_end and f_end > c_end and pd.isna(source_flow): 
                t_verdict = "EXPECTING_AMENDMENT"
                t_trace = "RULE_1: FED_END_DATE > CAY_END_DATE (DIRECT_FED_AMENDMENT_PENDING)"
            elif o_end and c_end and o_end > c_end: 
                t_verdict = "SEQUENCE_ANOMALY"
                t_trace = "RULE_2: ORACLE_END_DATE > CAYUSE_END_DATE (CRITICAL_PROCEDURAL_BREACH)"
            elif c_end and f_end and c_end > f_end:
                if (c_end - f_end).days <= 365:
                    t_verdict = "EXPANDED_AUTHORITIES_VALID"
                    t_trace = f"RULE_3A: CAY_END > FED_END BY {days_cayuse_minus_fed} DAYS (WITHIN_1_YEAR_EXPANDED_AUTHORITY)"
                else:
                    t_verdict = "EXPANDED_AUTHORITIES_EXCEEDED"
                    t_trace = f"RULE_3B: CAY_END > FED_END BY {days_cayuse_minus_fed} DAYS (EXCEEDS_1_YEAR_COMPLIANCE_MAX)"
            elif o_end != f_end or c_end != f_end: 
                t_verdict = "FEDERAL_REPORTING_LAG"
                t_trace = "RULE_4: SYSTEM_MISALIGNMENT_POTENTIAL_FEDERAL_LOG_LAG"
        if is_federal_funding and o_end and c_end and o_end > c_end: 
            t_verdict = "SEQUENCE_ANOMALY"
            t_trace = "RULE_2_ALT: INT_ORACLE_END > CAYUSE_END_ON_FED_PROJECT"
    elif o_end != c_end:
        if o_end and c_end and o_end > c_end and (o_end - c_end).days <= 90:
            t_verdict = "NON_FED_ADMIN_HOLD"
            t_trace = "RULE_5: ORACLE > CAYUSE BY <= 90 DAYS (NON_FED_CLOSEOUT_ADMIN_HOLD)"
        else:
            t_verdict = "INTERNAL_TIMELINE_MISALIGNMENT"
            t_trace = "RULE_6: ORACLE_END != CAYUSE_END (NO_FEDERAL_RECORD_FOUND)"

    if t_verdict in ["SEQUENCE_ANOMALY", "EXPANDED_AUTHORITIES_EXCEEDED"]: target_end_date = None  
    elif t_verdict == "EXPECTING_AMENDMENT": target_end_date = c_end 
    elif c_end and o_end and c_end > o_end: target_end_date = c_end 
    elif f_end and pd.isna(source_flow):
        dates_list = [d for d in [o_end, c_end, f_end] if d is not None]
        target_end_date = max(dates_list) if dates_list else o_end
    else: target_end_date = o_end if o_end else c_end

    is_timeline_actionable = "ACTIVE" in status_clean or ("EXPIRED" in status_clean and o_end and (PIPELINE_RUN_DATE - o_end).days <= 548)
    attention_timeline = "YES" if (o_end != target_end_date or c_end != target_end_date) and is_timeline_actionable and t_verdict not in ["IN_SYNC", "EXPANDED_AUTHORITIES_VALID", "NON_FED_ADMIN_HOLD", "EXPECTING_AMENDMENT"] else "NO"

    # Financial Engine
    if fed_cat == "UNMATCHED":
        if delta_oracle_limit_vs_cayuse > 1.00: 
            f_verdict = "INTERNAL_BUDGET_MISALIGNMENT"
            f_trace = f"RULE_7: ABS(ORACLE_LIMIT - CAYUSE) = ${delta_oracle_limit_vs_cayuse:,.2f} (INTERNAL_LEDGER_DRIFT)"
    else:
        if "CONTRACT" in source_mechanism.upper() or fed_cat in ["Prime Contract", "Subaward Contract"]:
            if abs(o_project_budget - fed_potential) > 1.00 or abs(o_header_limit - fed_potential) > 1.00: 
                f_verdict = "CONTRACT_VALUE_CEILING_MISMATCH"
                f_trace = "RULE_8: ORACLE_FINANCIALS_OUT_OF_BOUNDS_VS_FED_POTENTIAL_CEILING"
        else:
            if delta_oracle_limit_vs_cayuse > 1.00: 
                f_verdict = "INTERNAL_BUDGET_MISALIGNMENT"
                f_trace = f"RULE_9: ORACLE_LIMIT != CAYUSE BUDGET BY ${delta_oracle_limit_vs_cayuse:,.2f} (INTERNAL_LEDGER_DRIFT)"
            elif fed_potential < o_header_limit and delta_oracle_limit_vs_cayuse <= 1.00: 
                f_verdict = "FEDERAL_INCREMENTAL_CEILING_LAG"
                f_trace = "RULE_10: FED_POTENTIAL < ORACLE_HEADER_LIMIT (FEDERAL_INCREMENTAL_AWARD_LAG)"
            elif o_project_budget < fed_obligated: 
                f_verdict = "ORACLE_PROJECT_UNDER_FUNDED"
                f_trace = f"RULE_11A: ORACLE_PROJECT < FED_OBLIGATED BY ${abs(o_project_budget - fed_obligated):,.2f} (ORACLE_UNDER_ALLOCATED)"
            elif o_project_budget > fed_obligated: 
                f_verdict = "ORACLE_PROJECT_OVER_FUNDED"
                f_trace = f"RULE_11B: ORACLE_PROJECT > FED_OBLIGATED BY ${abs(o_project_budget - fed_obligated):,.2f} (ORACLE_OVER_ALLOCATED)"

    if fed_cat == "UNMATCHED": target_hard_limit = c_funding
    else: target_hard_limit = o_header_limit if (fed_potential < o_header_limit and delta_oracle_limit_vs_cayuse <= 1.00) else max(c_funding, fed_potential)
    target_project_budget = fed_obligated if fed_cat != "UNMATCHED" else c_funding

    attention_financial = "YES" if (abs(o_header_limit - target_hard_limit) > 1.00 or abs(o_project_budget - target_project_budget) > 1.00 or abs(c_funding - target_hard_limit) > 1.00) else "NO"
    if fed_cat == "UNMATCHED" and is_federal_funding: g_verdict = "UNTRACKED_FEDERAL_AWARD"

    oh_less, oh_more = np.nan, np.nan
    op_less, op_more = np.nan, np.nan
    cay_less, cay_more = np.nan, np.nan
    o_date_retract, o_date_extend = None, None
    c_date_retract, c_date_extend = None, None

    if t_verdict not in ["SEQUENCE_ANOMALY", "EXPANDED_AUTHORITIES_EXCEEDED", "EXPECTING_AMENDMENT"]:
        h_delta = target_hard_limit - o_header_limit
        p_delta = target_project_budget - o_project_budget
        c_delta = target_hard_limit - c_funding
        if h_delta < -1.00: oh_less = abs(h_delta)
        if h_delta > 1.00: oh_more = h_delta
        if p_delta < -1.00: op_less = abs(p_delta)
        if p_delta > 1.00: op_more = p_delta
        if c_delta < -1.00: cay_less = abs(c_delta)
        if c_delta > 1.00: cay_more = c_delta
        if target_end_date:
            if o_end and target_end_date < o_end: o_date_retract = target_end_date
            if o_end and target_end_date > o_end: o_date_extend = target_end_date
            if c_end and target_end_date < c_end: c_date_retract = target_end_date
            if c_end and target_end_date > c_end: c_date_extend = target_end_date

    if t_verdict in ["SEQUENCE_ANOMALY", "EXPANDED_AUTHORITIES_EXCEEDED"] or f_verdict == "CONTRACT_VALUE_CEILING_MISMATCH" or (o_header_limit > c_funding and is_federal_funding):
        route_class = "MANUAL_OR_AI_AUDIT_REQUIRED"
    elif attention_timeline == "YES" or attention_financial == "YES" or t_verdict == "EXPECTING_AMENDMENT":
        route_class = "AUTOMATED_CHANGE_READY"
    else:
        route_class = "SYSTEMS_FULLY_ALIGNED"

    triaged_records.append({
        'AWARD_NUMBER': clean_award_number,
        'FEDERAL_AWARD_IDENTIFIER': fed_id if fed_cat != "UNMATCHED" else "",
        'SPONSOR': row['Funding Source Name'],
        'AWARD_STATUS': source_status,
        'RECORD_RECONCILIATION_ROUTE': route_class,
        
        'DAYS_CAYUSE_MINUS_FED': days_cayuse_minus_fed,
        'DELTA_ORACLE_LIMIT_VS_CAYUSE': delta_oracle_limit_vs_cayuse,
        'DELTA_PROJECT_VS_FED_OBLIGATED': delta_project_vs_fed_obligated,
        'TIMELINE_RULE_TRACE': t_trace,
        'FINANCIAL_RULE_TRACE': f_trace,
        
        'ATTENTION_REQUIRED_TIMELINE': attention_timeline,
        'ATTENTION_REQUIRED_FINANCIAL': attention_financial,
        'TIMELINE_SYNC_STATUS': t_verdict,
        'FINANCIAL_SYNC_STATUS': f_verdict,
        'EXISTING_HARD_LIMIT_AMOUNT': o_header_limit,
        'EXISTING_BUDGET_BRDND_COST': o_project_budget,
        'EXISTING_CAYUSE_BUDGET': c_funding,
        'EXISTING_ORACLE_END_DATE': o_end,
        'EXISTING_CAYUSE_END_DATE': c_end,
        'FEDERAL_END_DATE': f_end,
        'PROPOSED_END_DATE': target_end_date,
        'PROPOSED_HARD_LIMIT_AMOUNT': target_hard_limit,
        'PROPOSED_BUDGET_BRDND_COST': target_project_budget,
        'HARD_LIMIT_AMOUNT_DECREASE': oh_less,
        'HARD_LIMIT_AMOUNT_INCREASE': oh_more,
        'BUDGET_BRDND_COST_DECREASE': op_less,
        'BUDGET_BRDND_COST_INCREASE': op_more,
        'CAYUSE_BUDGET_DECREASE': cay_less,
        'CAYUSE_BUDGET_INCREASE': cay_more,
        'ORACLE_DATE_RETRACT': o_date_retract,
        'ORACLE_DATE_EXTENSION': o_date_extend,
        'CAYUSE_DATE_RETRACT': c_date_retract,
        'CAYUSE_DATE_EXTENSION': c_date_extend
    })

full_df = pd.DataFrame(triaged_records)

# Export structured tables
with pd.ExcelWriter("Ecosystem_True_Up_Audit.xlsx", engine="openpyxl") as excel_writer:
    exec_headers = [
        'AWARD_NUMBER', 'FEDERAL_AWARD_IDENTIFIER', 'SPONSOR', 'AWARD_STATUS', 'RECORD_RECONCILIATION_ROUTE', 
        'DAYS_CAYUSE_MINUS_FED', 'DELTA_ORACLE_LIMIT_VS_CAYUSE', 'DELTA_PROJECT_VS_FED_OBLIGATED',
        'TIMELINE_RULE_TRACE', 'FINANCIAL_RULE_TRACE', 'ATTENTION_REQUIRED_TIMELINE', 'ATTENTION_REQUIRED_FINANCIAL',
        'EXISTING_ORACLE_END_DATE', 'EXISTING_CAYUSE_END_DATE', 'FEDERAL_END_DATE', 'PROPOSED_END_DATE',
        'EXISTING_HARD_LIMIT_AMOUNT', 'EXISTING_BUDGET_BRDND_COST', 'EXISTING_CAYUSE_BUDGET',
        'PROPOSED_HARD_LIMIT_AMOUNT', 'PROPOSED_BUDGET_BRDND_COST'
    ]
    save_as_structured_table(excel_writer, full_df[exec_headers], "Full Triage Master")
    
    time_headers = ['AWARD_NUMBER', 'FEDERAL_AWARD_IDENTIFIER', 'SPONSOR', 'DAYS_CAYUSE_MINUS_FED', 'TIMELINE_RULE_TRACE', 'EXISTING_ORACLE_END_DATE', 'EXISTING_CAYUSE_END_DATE', 'FEDERAL_END_DATE', 'PROPOSED_END_DATE', 'ORACLE_DATE_EXTENSION', 'ORACLE_DATE_RETRACT']
    save_as_structured_table(excel_writer, full_df[full_df['ATTENTION_REQUIRED_TIMELINE'] == 'YES'][time_headers], "Chronological Triage")
    
    fin_headers = ['AWARD_NUMBER', 'FEDERAL_AWARD_IDENTIFIER', 'SPONSOR', 'DELTA_ORACLE_LIMIT_VS_CAYUSE', 'DELTA_PROJECT_VS_FED_OBLIGATED', 'FINANCIAL_RULE_TRACE', 'EXISTING_HARD_LIMIT_AMOUNT', 'EXISTING_BUDGET_BRDND_COST', 'EXISTING_CAYUSE_BUDGET', 'PROPOSED_HARD_LIMIT_AMOUNT', 'PROPOSED_BUDGET_BRDND_COST', 'HARD_LIMIT_AMOUNT_INCREASE', 'HARD_LIMIT_AMOUNT_DECREASE', 'BUDGET_BRDND_COST_INCREASE', 'BUDGET_BRDND_COST_DECREASE']
    save_as_structured_table(excel_writer, full_df[full_df['ATTENTION_REQUIRED_FINANCIAL'] == 'YES'][fin_headers], "Financial Triage")

print("\n🎉 Analysis complete. Master verification workbook compiled into 'Ecosystem_True_Up_Audit.xlsx'!")