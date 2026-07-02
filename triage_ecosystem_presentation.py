import os
import pandas as pd
import numpy as np
import re
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

# File Input Environment Definitions
MASTER_INPUT = "Cayuse-Oracle award True Up - jth6-Copy.xlsx"
PRIME_ASST_FILE = "Assistance_PrimeAwardSummaries_2026-06-25_H19M12S13_1.csv"
PRIME_CONT_FILE = "Contracts_PrimeAwardSummaries_2026-06-25_H19M12S08_1.csv"
SUB_ASST_FILE = "Assistance_Subawards_2026-06-25_H19M12S41_1.csv"
SUB_CONT_FILE = "Contracts_Subawards_2026-06-25_H19M12S29_1.csv"

print("🚀 Launching Fact-Based Presentation Data Compiler (Ceiling Omitted)...")

if not os.path.exists(MASTER_INPUT):
    print(f"❌ Error: Cannot find master workbook '{MASTER_INPUT}' in this folder.")
    exit()

def advanced_clean_id(text):
    if pd.isna(text): return ""
    text = str(text).upper().strip()
    text = re.sub(r'^(NIH|NSF|DOD|DOHA|ONR|DOE|USDA|DHHS)\s+', '', text)
    text = re.sub(r'[^A-Z0-9-]', '', text)
    text = re.sub(r'-[A-Z0-9]{1,2}$', '', text) 
    text = re.sub(r'[^A-Z0-9]', '', text)
    if re.match(r'^[A-Z]+\d{7}$', text): text = re.sub(r'^[A-Z]+', '', text)
    if re.match(r'^\d[A-Z]\d{2}', text): text = text[1:]
    return text

def tokenize_string(s):
    if pd.isna(s): return []
    tokens = re.findall(r'[A-Z0-9\-]{4,}', str(s).upper())
    noise = {'RICE', 'WEST', 'GRANT', 'FEDERAL', 'CONTR', 'CONTRACT', 'AMENDMENT', 'MODIFICATION', 'BLANK', 'NONE', 'FOUNDATION', 'UNIVERSITY'}
    return [t for t in tokens if t not in noise]

def parse_mixed_date(val):
    if pd.isna(val): return None
    val_str = str(val).strip()
    if val_str.replace('.0', '').isdigit():
        clean_int = int(val_str.replace('.0', ''))
        return pd.to_datetime(clean_int, unit='D', origin='1899-12-30').date()
    return pd.to_datetime(val_str, errors='coerce').date()

def save_as_structured_table(writer, dataframe, name_of_sheet):
    dataframe.to_excel(writer, sheet_name=name_of_sheet, index=False)
    worksheet = writer.sheets[name_of_sheet]
    total_rows = len(dataframe) + 1
    total_cols = len(dataframe.columns) if len(dataframe) > 0 else 1
    column_letter = get_column_letter(total_cols)
    cell_range = f"A1:{column_letter}{total_rows}"
    table_id = "".join(c for c in name_of_sheet if c.isalnum())
    excel_table = Table(displayName=table_id, ref=cell_range)
    excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    worksheet.add_table(excel_table)

# Ingest Portfolios
print("📥 Ingesting local portfolio files and federal logs...")
master_df = pd.read_excel(MASTER_INPUT, sheet_name="Intermediate DATA")
gsum_df = pd.read_excel(MASTER_INPUT, sheet_name="Grant Summary")

gsum_df['clean_id'] = gsum_df['AWARD_NUMBER_SCRUB'].astype(str).str.upper().str.strip()
type_map = dict(zip(gsum_df['clean_id'], gsum_df['AWARD_TYPE_NAME']))
status_map = dict(zip(gsum_df['clean_id'], gsum_df['AWARD_STATUS']))
flow_through_map = dict(zip(gsum_df['clean_id'], gsum_df['FLOW_THROUGH_SPONSOR']))
header_limit_map = dict(zip(gsum_df['clean_id'], gsum_df['HARD_LIMIT_AMOUNT']))
project_budget_map = dict(zip(gsum_df['clean_id'], gsum_df['BUDGET_BRDND_COST']))

pa_df = pd.read_csv(PRIME_ASST_FILE, low_memory=False) if os.path.exists(PRIME_ASST_FILE) else pd.DataFrame(columns=['award_id_fain'])
pc_df = pd.read_csv(PRIME_CONT_FILE, low_memory=False) if os.path.exists(PRIME_CONT_FILE) else pd.DataFrame(columns=['award_id_piid'])
sa_df = pd.read_csv(SUB_ASST_FILE, low_memory=False) if os.path.exists(SUB_ASST_FILE) else pd.DataFrame(columns=['prime_award_fain'])
sc_df = pd.read_csv(SUB_CONT_FILE, low_memory=False) if os.path.exists(SUB_CONT_FILE) else pd.DataFrame(columns=['prime_award_piid'])

pa_lookup = {advanced_clean_id(r['award_id_fain']): (r['total_obligated_amount'], r['total_funding_amount'], r['period_of_performance_current_end_date'], r['award_id_fain'], "Prime Assistance") for _, r in pa_df.iterrows() if pd.notna(r['award_id_fain'])}
pc_lookup = {advanced_clean_id(r['award_id_piid']): (r['total_obligated_amount'], r['potential_total_value_of_award'], r['period_of_performance_current_end_date'], r['award_id_piid'], "Prime Contract") for _, r in pc_df.iterrows() if pd.notna(r['award_id_piid'])}

sa_lookup = {advanced_clean_id(r['prime_award_fain']): (r['subaward_amount'], r['subaward_amount'], r['prime_award_period_of_performance_current_end_date'], r['prime_award_fain'], "Subaward Assistance") for _, r in sa_df.iterrows() if pd.notna(r['prime_award_fain'])}
sc_lookup = {advanced_clean_id(r['prime_award_piid']): (r['subaward_amount'], r['subaward_amount'], r['prime_award_period_of_performance_current_end_date'], r['prime_award_piid'], "Subaward Contract") for _, r in sc_df.iterrows() if pd.notna(r['prime_award_piid'])}

fact_records = []

print("📊 Extracting system data coordinates into unvarnished data frames...")
for idx, row in master_df.iterrows():
    fed_id, fed_end_date, fed_obligated, fed_potential, fed_cat = np.nan, np.nan, 0.0, 0.0, "UNMATCHED"
    clean_award_number = str(row['Award Number']).replace('.0', '').upper().strip()
    
    source_status = status_map.get(clean_award_number, row['STATUS'] if 'STATUS' in master_df.columns else "Active")
    status_clean = str(source_status).upper().strip()
    if status_clean == "CLOSED": continue  

    proj_cayuse_val = str(row['Proj CAYUSE']).strip() if ('Proj CAYUSE' in master_df.columns and pd.notna(row['Proj CAYUSE'])) else ""

    primary_id_key = advanced_clean_id(row['Award Number'])
    name_tokens = tokenize_string(row['Award Name'])
    candidates = [primary_id_key] if primary_id_key else []
    candidates.extend([advanced_clean_id(t) for t in name_tokens if t])
    candidates = [c for c in candidates if c]
    
    for k in candidates:
        if k in pa_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = pa_lookup[k]; break
        elif k in pc_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = pc_lookup[k]; break
        elif k in sa_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = sa_lookup[k]; break
        elif k in sc_lookup: fed_obligated, fed_potential, fed_end_date, fed_id, fed_cat = sc_lookup[k]; break

    o_end = parse_mixed_date(row['OR END DATE'])
    c_end = parse_mixed_date(row['CAY END DATE'])
    f_end = parse_mixed_date(fed_end_date)
    
    c_funding = float(row['CAYUSE BUDGET']) if pd.notna(row['CAYUSE BUDGET']) else 0.0
    o_header_limit = float(header_limit_map.get(clean_award_number, c_funding))
    o_project_budget = float(project_budget_map.get(clean_award_number, float(row['OR Allocated Funding']) if pd.notna(row['OR Allocated Funding']) else 0.0))
    source_type = type_map.get(clean_award_number, row['Award Type'] if 'Award Type' in master_df.columns else "Federal")

    fact_records.append({
        # System Identifiers
        'ORACLE_AWARD_NUMBER': clean_award_number,
        'CAYUSE_PROJECT_NUMBER': proj_cayuse_val if (proj_cayuse_val.upper() != "NOT FOUND" and proj_cayuse_val.upper() != "NAN") else "N/A",
        'FEDERAL_AWARD_IDENTIFIER': fed_id if fed_cat != "UNMATCHED" else "N/A",
        'SPONSOR_NAME': row['Funding Source Name'],
        'SPONSOR_TYPE': source_type,
        'AWARD_STATUS': source_status,
        
        # Absolute End Dates Group
        'ORACLE_END_DATE': o_end,
        'CAYUSE_END_DATE': c_end,
        'FEDERAL_END_DATE': f_end,
        
        # Budget Ledger Values Group (Potential Ceiling completely removed)
        'ORACLE_HEADER_HARD_LIMIT': o_header_limit,
        'ORACLE_PROJECT_ALLOCATED_BUDGET': o_project_budget,
        'CAYUSE_TOTAL_AUTHORIZED_BUDGET': c_funding,
        'FEDERAL_USA_OBLIGATED_AMOUNT': fed_obligated if fed_cat != "UNMATCHED" else 0.0
    })

full_df = pd.DataFrame(fact_records)

# Generate the three targeted presentation sheets
OUTPUT_FILE = "Sponsor_Ecosystem_Fact_Baseline.xlsx"
with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    
    # Tab 1: Comprehensive Master Baseline
    master_cols = [
        'ORACLE_AWARD_NUMBER', 'CAYUSE_PROJECT_NUMBER', 'FEDERAL_AWARD_IDENTIFIER', 'SPONSOR_NAME', 'SPONSOR_TYPE', 'AWARD_STATUS',
        'ORACLE_END_DATE', 'CAYUSE_END_DATE', 'FEDERAL_END_DATE',
        'ORACLE_HEADER_HARD_LIMIT', 'ORACLE_PROJECT_ALLOCATED_BUDGET', 'CAYUSE_TOTAL_AUTHORIZED_BUDGET', 'FEDERAL_USA_OBLIGATED_AMOUNT'
    ]
    save_as_structured_table(writer, full_df[master_cols], "Master Baseline Sheet")
    
    # Tab 2: Chronological Baseline
    time_cols = [
        'ORACLE_AWARD_NUMBER', 'CAYUSE_PROJECT_NUMBER', 'FEDERAL_AWARD_IDENTIFIER', 'SPONSOR_NAME', 'SPONSOR_TYPE', 'AWARD_STATUS',
        'ORACLE_END_DATE', 'CAYUSE_END_DATE', 'FEDERAL_END_DATE'
    ]
    save_as_structured_table(writer, full_df[time_cols], "Chronological Baseline")
    
    # Tab 3: Financial Baseline
    fin_cols = [
        'ORACLE_AWARD_NUMBER', 'CAYUSE_PROJECT_NUMBER', 'FEDERAL_AWARD_IDENTIFIER', 'SPONSOR_NAME', 'SPONSOR_TYPE', 'AWARD_STATUS',
        'ORACLE_HEADER_HARD_LIMIT', 'ORACLE_PROJECT_ALLOCATED_BUDGET', 'CAYUSE_TOTAL_AUTHORIZED_BUDGET', 'FEDERAL_USA_OBLIGATED_AMOUNT'
    ]
    save_as_structured_table(writer, full_df[fin_cols], "Financial Baseline")

print(f"\n🎉 Success! The presentation workbook has been compiled at '{OUTPUT_FILE}'!")